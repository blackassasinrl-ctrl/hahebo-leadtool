"""
Genereert HAHEBO_leadbestand_template.xlsx

Bouwt alleen het lege Excel-sjabloon voor het toekomstige leadbestand
(kolomstructuur, validatie, opmaak, uitlegtabblad). Geen leadgeneratie,
geen scraping, geen internet, geen mailingfunctionaliteit.
"""

from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
import os

OUTPUT_PATH = os.path.join(os.path.dirname(__file__), "HAHEBO_leadbestand_template.xlsx")

COLUMNS = [
    ("Lead ID", 12),
    ("Bedrijfsnaam", 32),
    ("KVK-nummer", 14),
    ("Domein", 22),
    ("Website", 30),
    ("Vestigingsplaats", 20),
    ("Postcode", 12),
    ("Adres", 30),
    ("Aantal medewerkers", 18),
    ("Medewerkersklasse", 18),
    ("Bron medewerkers", 20),
    ("Bronniveau medewerkers", 18),
    ("Peildatum medewerkers", 18),
    ("Actualiteit medewerkers", 20),
    ("Meetniveau medewerkers", 18),
    ("Eenheid medewerkers", 18),
    ("Handmatig beoordeeld", 16),
    ("Bronnen handmatige beoordeling", 32),
    ("Segment", 16),
    ("Doelgroepstatus", 20),
    ("Algemeen e-mailadres", 28),
    ("Telefoonnummer", 16),
    ("Bron contactgegevens", 20),
    ("Resultaat vergelijking acquisitielijst", 26),
    ("Matchwijze acquisitielijst", 22),
    ("Datum verzameld", 16),
    ("Opmerkingen", 30),
]

SEGMENT_OPTIONS = ["50+", "regionaal <50", "nog te bepalen", "onbekend"]
MATCH_RESULT_OPTIONS = ["nog niet gecontroleerd", "nieuw", "bestaat al"]
BRONNIVEAU_OPTIONS = ["A", "B", "C", "D", "Onbekend"]
ACTUALITEIT_OPTIONS = ["Expliciet gedateerd", "Actueel/live", "Historisch/verouderd", "Actualiteit onduidelijk"]
MEETNIVEAU_OPTIONS = ["Organisatie", "Entiteit", "Vestiging", "Onbekend"]
EENHEID_OPTIONS = ["Werkzame personen", "FTE", "Onbekend"]
HANDMATIG_BEOORDEELD_OPTIONS = ["Ja", "Nee"]
DOELGROEPSTATUS_OPTIONS = ["binnen doelgroep", "buiten doelgroep", "nog niet te bepalen"]

COLUMN_EXPLANATIONS = {
    "Lead ID": "Uniek volgnummer of code voor deze lead, bijvoorbeeld TEST001.",
    "Bedrijfsnaam": "Officiele handelsnaam van het bedrijf.",
    "KVK-nummer": "KVK-nummer als tekst, zodat voorloopnullen behouden blijven.",
    "Domein": "Domeinnaam van het bedrijf, bijvoorbeeld voorbeeldbedrijf.nl. 'Niet gevonden' als er geen betrouwbare website is gevonden.",
    "Website": "Volledige website-URL van het bedrijf. 'Niet gevonden' als er geen betrouwbare website is gevonden.",
    "Vestigingsplaats": "Plaats waar het bedrijf is gevestigd.",
    "Postcode": "Postcode als tekst, zodat de notatie (incl. voorloopnullen) behouden blijft.",
    "Adres": "Straatnaam en huisnummer van het vestigingsadres.",
    "Aantal medewerkers": "Het gekozen medewerkersgegeven dat voor de beoordeling wordt gebruikt. Nooit zelf verzinnen of berekenen. Eenheid medewerkers geeft aan of dit werkzame personen, FTE of onbekend betreft.",
    "Medewerkersklasse": "Categorie waarin het aantal medewerkers valt, bijvoorbeeld een grootteklasse.",
    "Bron medewerkers": "Bron waaruit het aantal medewerkers afkomstig is.",
    "Bronniveau medewerkers": "Keuzelijst: A / B / C / D / Onbekend. Geeft het kwaliteitsniveau van de bron voor het medewerkersaantal aan.",
    "Peildatum medewerkers": "Datum waarop het medewerkersaantal is gemeten, notatie dd-mm-jjjj. Tekst 'Onbekend' als geen peildatum kan worden vastgesteld.",
    "Actualiteit medewerkers": (
        "Keuzelijst: Expliciet gedateerd / Actueel/live / Historisch/verouderd / Actualiteit onduidelijk. "
        "Expliciet gedateerd: de bron noemt een concrete peildatum, boekjaar of datum waarop het medewerkersgegeven betrekking heeft. "
        "Actueel/live: de bron presenteert het gegeven als huidige situatie op een actuele pagina, maar noemt geen expliciete peildatum. "
        "Historisch/verouderd: het gegeven heeft duidelijk betrekking op een afgesloten of verouderde periode. "
        "Actualiteit onduidelijk: niet betrouwbaar vast te stellen of het gegeven nog actueel is."
    ),
    "Meetniveau medewerkers": (
        "Keuzelijst: Organisatie / Entiteit / Vestiging / Onbekend. "
        "Organisatie: de onderneming of organisatie als commercieel geheel waarop HAHEBO de prospectbenadering richt. "
        "Entiteit: een juridische entiteit, bijvoorbeeld een B.V. binnen een grotere organisatie. "
        "Vestiging: een fysieke locatie. "
        "Onbekend: niet betrouwbaar vast te stellen op welk niveau het medewerkersgegeven betrekking heeft."
    ),
    "Eenheid medewerkers": "Keuzelijst: Werkzame personen / FTE / Onbekend. Werkzame personen is de voorkeursmaatstaf voor de grens van 50 medewerkers.",
    "Handmatig beoordeeld": "Keuzelijst: Ja / Nee. Geeft aan of deze lead handmatig is beoordeeld.",
    "Bronnen handmatige beoordeling": "Vrij tekstveld voor een of meerdere gebruikte bron-URL's of een korte verwijzing naar de gebruikte bronnen. Mag leeg blijven als Handmatig beoordeeld = Nee.",
    "Segment": "Keuzelijst: 50+ / regionaal <50 / nog te bepalen / onbekend.",
    "Doelgroepstatus": (
        "Keuzelijst: binnen doelgroep / buiten doelgroep / nog niet te bepalen. "
        "Geeft aan of een lead binnen de commerciele doelgroep van HAHEBO valt nadat onder andere geografische criteria zijn toegepast. "
        "binnen doelgroep: lead voldoet aan de vastgestelde doelgroepcriteria. "
        "buiten doelgroep: lead is wel geidentificeerd maar valt op basis van vastgestelde criteria buiten de doelgroep. "
        "nog niet te bepalen: nog niet alle criteria zijn door HAHEBO vastgesteld of voldoende informatie ontbreekt."
    ),
    "Algemeen e-mailadres": "Algemeen of centraal e-mailadres van het bedrijf.",
    "Telefoonnummer": "Telefoonnummer als tekst, zodat voorloopnullen behouden blijven.",
    "Bron contactgegevens": "Bron waaruit het e-mailadres en/of telefoonnummer afkomstig is.",
    "Resultaat vergelijking acquisitielijst": "Keuzelijst: nog niet gecontroleerd / nieuw / bestaat al. Standaard 'nog niet gecontroleerd' bij nieuwe regels.",
    "Matchwijze acquisitielijst": "Manier waarop de vergelijking met de acquisitielijst is uitgevoerd (bijv. op naam, KVK-nummer of domein).",
    "Datum verzameld": "Datum waarop de leadgegevens zijn verzameld, notatie dd-mm-jjjj.",
    "Opmerkingen": "Vrij invulveld voor aanvullende opmerkingen of bijzonderheden.",
}

TEXT_COLUMNS = {"KVK-nummer", "Telefoonnummer", "Postcode"}

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def build_workbook():
    wb = Workbook()

    ws = wb.active
    ws.title = "Leadbestand"

    # Header row
    for col_idx, (name, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=name)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}1"

    col_index = {name: i + 1 for i, (name, _) in enumerate(COLUMNS)}

    # Force text format on KVK-nummer, Telefoonnummer, Postcode for a large range of rows
    text_format_rows = 500
    for col_name in TEXT_COLUMNS:
        c = col_index[col_name]
        for r in range(2, text_format_rows + 1):
            ws.cell(row=r, column=c).number_format = "@"

    # Date format on "Datum verzameld" and "Peildatum medewerkers"
    # (Peildatum medewerkers may also contain the text "Onbekend")
    for date_col_name in ("Datum verzameld", "Peildatum medewerkers"):
        date_col = col_index[date_col_name]
        for r in range(2, text_format_rows + 1):
            ws.cell(row=r, column=date_col).number_format = "dd-mm-yyyy"

    # Data validation: Segment
    segment_col_letter = get_column_letter(col_index["Segment"])
    dv_segment = DataValidation(
        type="list",
        formula1='"' + ",".join(SEGMENT_OPTIONS) + '"',
        allow_blank=True,
        showDropDown=False,
    )
    dv_segment.error = "Kies een geldige waarde uit de lijst."
    dv_segment.errorTitle = "Ongeldige invoer"
    dv_segment.prompt = "Kies een segment uit de lijst."
    dv_segment.promptTitle = "Segment"
    ws.add_data_validation(dv_segment)
    dv_segment.add(f"{segment_col_letter}2:{segment_col_letter}{text_format_rows}")

    # Data validation: Resultaat vergelijking acquisitielijst
    match_col_letter = get_column_letter(col_index["Resultaat vergelijking acquisitielijst"])
    dv_match = DataValidation(
        type="list",
        formula1='"' + ",".join(MATCH_RESULT_OPTIONS) + '"',
        allow_blank=True,
        showDropDown=False,
    )
    dv_match.error = "Kies een geldige waarde uit de lijst."
    dv_match.errorTitle = "Ongeldige invoer"
    dv_match.prompt = "Kies het resultaat van de vergelijking met de acquisitielijst."
    dv_match.promptTitle = "Resultaat vergelijking"
    ws.add_data_validation(dv_match)
    dv_match.add(f"{match_col_letter}2:{match_col_letter}{text_format_rows}")

    # Data validation: Bronniveau medewerkers
    bronniveau_col_letter = get_column_letter(col_index["Bronniveau medewerkers"])
    dv_bronniveau = DataValidation(
        type="list",
        formula1='"' + ",".join(BRONNIVEAU_OPTIONS) + '"',
        allow_blank=True,
        showDropDown=False,
    )
    dv_bronniveau.error = "Kies een geldige waarde uit de lijst."
    dv_bronniveau.errorTitle = "Ongeldige invoer"
    dv_bronniveau.prompt = "Kies het bronniveau voor het medewerkersaantal."
    dv_bronniveau.promptTitle = "Bronniveau medewerkers"
    ws.add_data_validation(dv_bronniveau)
    dv_bronniveau.add(f"{bronniveau_col_letter}2:{bronniveau_col_letter}{text_format_rows}")

    # Data validation: Actualiteit medewerkers
    actualiteit_col_letter = get_column_letter(col_index["Actualiteit medewerkers"])
    dv_actualiteit = DataValidation(
        type="list",
        formula1='"' + ",".join(ACTUALITEIT_OPTIONS) + '"',
        allow_blank=True,
        showDropDown=False,
    )
    dv_actualiteit.error = "Kies een geldige waarde uit de lijst."
    dv_actualiteit.errorTitle = "Ongeldige invoer"
    dv_actualiteit.prompt = "Kies: Expliciet gedateerd / Actueel/live / Historisch/verouderd / Actualiteit onduidelijk."
    dv_actualiteit.promptTitle = "Actualiteit medewerkers"
    ws.add_data_validation(dv_actualiteit)
    dv_actualiteit.add(f"{actualiteit_col_letter}2:{actualiteit_col_letter}{text_format_rows}")

    # Data validation: Meetniveau medewerkers
    meetniveau_col_letter = get_column_letter(col_index["Meetniveau medewerkers"])
    dv_meetniveau = DataValidation(
        type="list",
        formula1='"' + ",".join(MEETNIVEAU_OPTIONS) + '"',
        allow_blank=True,
        showDropDown=False,
    )
    dv_meetniveau.error = "Kies een geldige waarde uit de lijst."
    dv_meetniveau.errorTitle = "Ongeldige invoer"
    dv_meetniveau.prompt = "Kies het meetniveau van het medewerkersaantal."
    dv_meetniveau.promptTitle = "Meetniveau medewerkers"
    ws.add_data_validation(dv_meetniveau)
    dv_meetniveau.add(f"{meetniveau_col_letter}2:{meetniveau_col_letter}{text_format_rows}")

    # Data validation: Eenheid medewerkers
    eenheid_col_letter = get_column_letter(col_index["Eenheid medewerkers"])
    dv_eenheid = DataValidation(
        type="list",
        formula1='"' + ",".join(EENHEID_OPTIONS) + '"',
        allow_blank=True,
        showDropDown=False,
    )
    dv_eenheid.error = "Kies een geldige waarde uit de lijst."
    dv_eenheid.errorTitle = "Ongeldige invoer"
    dv_eenheid.prompt = "Werkzame personen is de voorkeursmaatstaf voor de grens van 50 medewerkers."
    dv_eenheid.promptTitle = "Eenheid medewerkers"
    ws.add_data_validation(dv_eenheid)
    dv_eenheid.add(f"{eenheid_col_letter}2:{eenheid_col_letter}{text_format_rows}")

    # Data validation: Handmatig beoordeeld
    handmatig_col_letter = get_column_letter(col_index["Handmatig beoordeeld"])
    dv_handmatig = DataValidation(
        type="list",
        formula1='"' + ",".join(HANDMATIG_BEOORDEELD_OPTIONS) + '"',
        allow_blank=True,
        showDropDown=False,
    )
    dv_handmatig.error = "Kies een geldige waarde uit de lijst."
    dv_handmatig.errorTitle = "Ongeldige invoer"
    dv_handmatig.prompt = "Is deze lead handmatig beoordeeld?"
    dv_handmatig.promptTitle = "Handmatig beoordeeld"
    ws.add_data_validation(dv_handmatig)
    dv_handmatig.add(f"{handmatig_col_letter}2:{handmatig_col_letter}{text_format_rows}")

    # Data validation: Doelgroepstatus
    doelgroepstatus_col_letter = get_column_letter(col_index["Doelgroepstatus"])
    dv_doelgroepstatus = DataValidation(
        type="list",
        formula1='"' + ",".join(DOELGROEPSTATUS_OPTIONS) + '"',
        allow_blank=True,
        showDropDown=False,
    )
    dv_doelgroepstatus.error = "Kies een geldige waarde uit de lijst."
    dv_doelgroepstatus.errorTitle = "Ongeldige invoer"
    dv_doelgroepstatus.prompt = "Kies: binnen doelgroep / buiten doelgroep / nog niet te bepalen."
    dv_doelgroepstatus.promptTitle = "Doelgroepstatus"
    ws.add_data_validation(dv_doelgroepstatus)
    dv_doelgroepstatus.add(f"{doelgroepstatus_col_letter}2:{doelgroepstatus_col_letter}{text_format_rows}")

    # Fictional example row
    example = {
        "Lead ID": "TEST001",
        "Bedrijfsnaam": "TESTDATA - verwijderen voor gebruik",
        "KVK-nummer": "01234567",
        "Domein": "voorbeeldbedrijf.nl",
        "Website": "https://www.voorbeeldbedrijf.nl",
        "Vestigingsplaats": "Voorbeeldstad",
        "Postcode": "1234 AB",
        "Adres": "Voorbeeldstraat 1",
        "Aantal medewerkers": 75,
        "Medewerkersklasse": "50+",
        "Bron medewerkers": "https://example.invalid/medewerkers",
        "Bronniveau medewerkers": "B",
        "Peildatum medewerkers": "01-03-2025",
        "Actualiteit medewerkers": "Expliciet gedateerd",
        "Meetniveau medewerkers": "Vestiging",
        "Eenheid medewerkers": "Werkzame personen",
        "Handmatig beoordeeld": "Ja",
        "Bronnen handmatige beoordeling": "https://example.invalid/handmatige-beoordeling",
        "Segment": "50+",
        "Doelgroepstatus": "nog niet te bepalen",
        "Algemeen e-mailadres": "info@voorbeeldbedrijf.nl",
        "Telefoonnummer": "0201234567",
        "Bron contactgegevens": "https://example.invalid/contact",
        "Resultaat vergelijking acquisitielijst": "nog niet gecontroleerd",
        "Matchwijze acquisitielijst": "",
        "Datum verzameld": "15-09-2026",
        "Opmerkingen": "Fictieve voorbeeldregel, uitsluitend ter illustratie van de kolomstructuur.",
    }
    for name, value in example.items():
        c = col_index[name]
        cell = ws.cell(row=2, column=c, value=(value if value != "" else None))
        if name in TEXT_COLUMNS:
            cell.number_format = "@"
        if name in ("Datum verzameld", "Peildatum medewerkers"):
            cell.number_format = "dd-mm-yyyy"

    # --- Sheet 2: Uitleg ---
    ws2 = wb.create_sheet("Uitleg")
    ws2.append(["Kolom", "Uitleg"])
    ws2["A1"].font = HEADER_FONT
    ws2["B1"].font = HEADER_FONT
    ws2["A1"].fill = HEADER_FILL
    ws2["B1"].fill = HEADER_FILL
    ws2.freeze_panes = "A2"
    ws2.auto_filter.ref = "A1:B1"

    for name, _ in COLUMNS:
        ws2.append([name, COLUMN_EXPLANATIONS[name]])

    ws2.column_dimensions["A"].width = 32
    ws2.column_dimensions["B"].width = 90
    for row in ws2.iter_rows(min_row=2, max_row=ws2.max_row, max_col=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    return wb


def main():
    wb = build_workbook()
    wb.save(OUTPUT_PATH)
    print(f"Bestand opgeslagen: {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
