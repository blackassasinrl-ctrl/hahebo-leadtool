"""
BRIN -> aantal personen lookup op basis van lokaal opgeslagen DUO-bestanden.

Gebruik:
    python scripts/duo_onderwijs_lookup.py <BRIN>
    python scripts/duo_onderwijs_lookup.py --selftest

Voorbeelden:
    python scripts/duo_onderwijs_lookup.py 17HN
    python scripts/duo_onderwijs_lookup.py 20GB

Bronbestanden (alleen lezen, nooit gewijzigd):
    data/duo/onderwijspersoneel_vo_personen_2011_2025.xlsx  (officieel DUO-bestand)
    data/duo/vestigingen_vo.csv                              (officieel DUO-bestand)
    data/duo/duo_po_brin_personen_2011_2025_compact.xlsx     (AFGELEIDE compacte dataset)
    data/duo/vestigingen_sbo_so_vso.csv                      (officieel DUO-bestand)

Werkwijze en beslisregels: zie data/duo/README.md.

Dit script bepaalt alleen per BRIN/instellingscode een sector, vestigingen en
(indien eenduidig) een actueel personeelsaantal. Het past geen leadbestanden
aan en legt geen automatische koppeling met het leadbestand.
"""

import csv
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

BASE_DIR = Path(__file__).resolve().parent.parent
DUO_DIR = BASE_DIR / "data" / "duo"

VO_PERSONEEL_XLSX = DUO_DIR / "onderwijspersoneel_vo_personen_2011_2025.xlsx"
VO_PERSONEEL_SHEET = "owtype-best-instelling"
VO_VESTIGINGEN_CSV = DUO_DIR / "vestigingen_vo.csv"

PO_PERSONEEL_XLSX = DUO_DIR / "duo_po_brin_personen_2011_2025_compact.xlsx"
SBO_SO_VSO_VESTIGINGEN_CSV = DUO_DIR / "vestigingen_sbo_so_vso.csv"

YEAR_COLUMN_RE = re.compile(r"^PERSONEN (\d{4})$")
NUMERIC_RE = re.compile(r"^-?\d+$")


# ---------------------------------------------------------------------------
# BRIN-normalisatie
# ---------------------------------------------------------------------------

def normalize_brin(raw):
    """Normaliseert BRIN-invoer: hoofdletterongevoelig, trimt spaties.

    Accepteert uitsluitend een code van exact 4 tekens. Een vestigingscode
    (zoals '17HN00') wordt NOOIT stilzwijgend ingekort tot een BRIN.
    """
    stripped = raw.strip()
    normalized = stripped.upper()
    if len(normalized) != 4:
        raise ValueError(
            f"Ongeldige invoer '{stripped}': een BRIN/instellingscode moet exact "
            f"4 tekens lang zijn (ontvangen: {len(normalized)} tekens). Een "
            f"vestigingscode zoals '17HN00' wordt niet automatisch omgezet naar "
            f"een 4-tekens BRIN. Geef uitsluitend de 4-tekens BRIN op (bijv. 17HN)."
        )
    return normalized


# ---------------------------------------------------------------------------
# Bevoegd gezag normalisatie
# ---------------------------------------------------------------------------

def normalize_bevoegd_gezag(value):
    """Normaliseert BEVOEGD GEZAG naar tekst zonder onnodige '.0'.

    Werkt voor int, float en tekst-representaties (bv. 41211, 41211.0,
    "41211", "41211.0" -> "41211").
    """
    if value is None:
        return None
    s = str(value).strip()
    if s == "":
        return None
    try:
        f = float(s)
    except ValueError:
        return s
    if f.is_integer():
        return str(int(f))
    return s


# ---------------------------------------------------------------------------
# CSV-vestigingsbestanden
# ---------------------------------------------------------------------------

def _read_csv_rows(path):
    with open(path, encoding="cp1252", newline="") as f:
        reader = csv.DictReader(f, delimiter=";")
        return list(reader)


def _vestiging_from_row(row):
    straat = (row.get("STRAATNAAM") or "").strip()
    huisnr = (row.get("HUISNUMMER-TOEVOEGING") or "").strip()
    adres = f"{straat} {huisnr}".strip()
    return {
        "vestigingscode": (row.get("VESTIGINGSCODE") or "").strip(),
        "vestigingsnaam": (row.get("VESTIGINGSNAAM") or "").strip(),
        "adres": adres,
        "postcode": (row.get("POSTCODE") or "").strip(),
        "plaats": (row.get("PLAATSNAAM") or "").strip(),
        "bevoegd_gezag_nummer": normalize_bevoegd_gezag(row.get("BEVOEGD GEZAG NUMMER")),
    }


def find_vestigingen(brin, csv_path):
    rows = _read_csv_rows(csv_path)
    matches = []
    for row in rows:
        code = (row.get("INSTELLINGSCODE") or "").strip().upper()
        if code == brin:
            matches.append(_vestiging_from_row(row))
    return matches


# ---------------------------------------------------------------------------
# XLSX-personeelsbestanden
# ---------------------------------------------------------------------------

def _read_xlsx_rows(path, sheet_name=None):
    """Leest een werkblad als lijst van dicts (header -> waarde), read-only."""
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[sheet_name] if sheet_name else wb.active
        rows_iter = ws.iter_rows(values_only=True)
        header = [str(h).strip() if h is not None else "" for h in next(rows_iter)]
        rows = []
        for values in rows_iter:
            if values is None or all(v is None for v in values):
                continue
            rows.append(dict(zip(header, values)))
        return header, rows
    finally:
        wb.close()


def detect_year_columns(header):
    """Vindt kolommen die exact voldoen aan 'PERSONEN YYYY'. Geeft {jaar: kolomnaam}."""
    year_cols = {}
    for name in header:
        m = YEAR_COLUMN_RE.match(name)
        if m:
            year_cols[int(m.group(1))] = name
    return year_cols


def find_personeel_rows(brin, xlsx_path, sheet_name=None):
    header, rows = _read_xlsx_rows(xlsx_path, sheet_name)
    year_cols = detect_year_columns(header)
    matches = []
    for row in rows:
        code = str(row.get("INSTELLINGSCODE") or "").strip().upper()
        if code == "BOVENSCHOOLS":
            continue
        if code == brin:
            matches.append(row)
    return matches, year_cols


def try_numeric(value):
    """Geeft int terug voor een echte numerieke waarde, anders None.

    '*', '**', '<5', lege cel en andere tekst worden NOOIT naar 0 omgezet.
    """
    if value is None:
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return int(value) if value.is_integer() else None
    s = str(value).strip()
    if NUMERIC_RE.match(s):
        return int(s)
    return None


def bepaal_actuele_waarde(row, year_cols):
    """Bepaalt actueel jaar/waarde plus historische numerieke context."""
    if not year_cols:
        return {
            "actueel_jaar": None,
            "actuele_bronwaarde": None,
            "aantal_personen": None,
            "waarde_status": None,
            "peildatum": None,
            "oorspronkelijke_waarde": None,
            "laatste_numerieke_waarde": None,
            "laatste_numerieke_jaar": None,
            "laatste_numerieke_peildatum": None,
            "segment": None,
        }

    jaren_aflopend = sorted(year_cols.keys(), reverse=True)
    actueel_jaar = jaren_aflopend[0]
    actuele_kolom = year_cols[actueel_jaar]
    actuele_bronwaarde = row.get(actuele_kolom)
    numeriek = try_numeric(actuele_bronwaarde)

    # Historische context: meest recente OUDERE jaar met een numerieke waarde.
    laatste_numerieke_waarde = None
    laatste_numerieke_jaar = None
    laatste_numerieke_peildatum = None
    for jaar in jaren_aflopend[1:]:
        kandidaat = try_numeric(row.get(year_cols[jaar]))
        if kandidaat is not None:
            laatste_numerieke_waarde = kandidaat
            laatste_numerieke_jaar = jaar
            laatste_numerieke_peildatum = f"01-10-{jaar}"
            break

    if numeriek is not None:
        segment = "50+" if numeriek >= 50 else "nog te bepalen"
        return {
            "actueel_jaar": actueel_jaar,
            "actuele_bronwaarde": actuele_bronwaarde,
            "aantal_personen": numeriek,
            "waarde_status": "numeriek",
            "peildatum": f"01-10-{actueel_jaar}",
            "oorspronkelijke_waarde": None,
            "laatste_numerieke_waarde": laatste_numerieke_waarde,
            "laatste_numerieke_jaar": laatste_numerieke_jaar,
            "laatste_numerieke_peildatum": laatste_numerieke_peildatum,
            "segment": segment,
        }

    return {
        "actueel_jaar": actueel_jaar,
        "actuele_bronwaarde": actuele_bronwaarde,
        "aantal_personen": None,
        "waarde_status": "niet_numeriek",
        "peildatum": None,
        "oorspronkelijke_waarde": actuele_bronwaarde,
        "laatste_numerieke_waarde": laatste_numerieke_waarde,
        "laatste_numerieke_jaar": laatste_numerieke_jaar,
        "laatste_numerieke_peildatum": laatste_numerieke_peildatum,
        "segment": "onbekend",
    }


# ---------------------------------------------------------------------------
# Hoofdlogica
# ---------------------------------------------------------------------------

def _leeg_resultaat(brin):
    return {
        "brin": brin,
        "sector": None,
        "status": None,
        "instellingsnaam_indicatie": None,
        "bevoegd_gezag_nummer": None,
        "aantal_vestigingen": 0,
        "vestigingen": [],
        "personeelsbron": None,
        "personeelsbron_type": None,
        "actueel_jaar": None,
        "actuele_bronwaarde": None,
        "aantal_personen": None,
        "waarde_status": None,
        "peildatum": None,
        "laatste_numerieke_waarde": None,
        "laatste_numerieke_jaar": None,
        "laatste_numerieke_peildatum": None,
        "bronniveau_medewerkers": None,
        "actualiteit_medewerkers": None,
        "meetniveau_medewerkers": None,
        "eenheid_medewerkers": None,
        "segment": None,
        "opmerkingen": [],
    }


def _instellingsnaam_indicatie(brin, vestigingen):
    if not vestigingen:
        return None
    hoofd = next((v for v in vestigingen if v["vestigingscode"] == f"{brin}00"), None)
    gekozen = hoofd or vestigingen[0]
    return gekozen["vestigingsnaam"] or None


def lookup_brin(raw_brin):
    brin = normalize_brin(raw_brin)
    result = _leeg_resultaat(brin)

    vo_vestigingen = find_vestigingen(brin, VO_VESTIGINGEN_CSV)
    po_vestigingen = find_vestigingen(brin, SBO_SO_VSO_VESTIGINGEN_CSV)

    in_vo = len(vo_vestigingen) > 0
    in_po = len(po_vestigingen) > 0

    if in_vo and in_po:
        result["status"] = "handmatige_controle"
        result["sector"] = "onduidelijk (komt voor in zowel vestigingen_vo.csv als vestigingen_sbo_so_vso.csv)"
        result["vestigingen"] = vo_vestigingen + po_vestigingen
        result["aantal_vestigingen"] = len(result["vestigingen"])
        result["instellingsnaam_indicatie"] = _instellingsnaam_indicatie(brin, result["vestigingen"])
        result["opmerkingen"].append(
            "BRIN komt voor in zowel het VO- als het SBO/SO/VSO-vestigingenbestand. "
            "Sector is niet automatisch bepaald; handmatige controle nodig. "
            "Er is geen personeelsbron geraadpleegd."
        )
        return result

    if not in_vo and not in_po:
        result["status"] = "niet_gevonden"
        result["opmerkingen"].append(
            "BRIN komt niet voor in vestigingen_vo.csv en niet in vestigingen_sbo_so_vso.csv."
        )
        return result

    if in_vo:
        result["sector"] = "VO"
        vestigingen = vo_vestigingen
        personeel_rows, year_cols = find_personeel_rows(brin, VO_PERSONEEL_XLSX, VO_PERSONEEL_SHEET)
        result["personeelsbron"] = f"data/duo/{VO_PERSONEEL_XLSX.name} (werkblad: {VO_PERSONEEL_SHEET})"
        result["personeelsbron_type"] = "officieel_DUO_bestand"
    else:
        result["sector"] = "PO/speciaal onderwijs"
        vestigingen = po_vestigingen
        personeel_rows, year_cols = find_personeel_rows(brin, PO_PERSONEEL_XLSX)
        result["personeelsbron"] = f"data/duo/{PO_PERSONEEL_XLSX.name}"
        result["personeelsbron_type"] = "afgeleide_compacte_DUO_dataset"
        result["opmerkingen"].append(
            "Let op: de PO-personeelsbron is een AFGELEIDE compacte DUO-dataset, "
            "niet het originele officiële DUO-bestand (zie data/duo/README.md)."
        )

    result["vestigingen"] = vestigingen
    result["aantal_vestigingen"] = len(vestigingen)
    result["instellingsnaam_indicatie"] = _instellingsnaam_indicatie(brin, vestigingen)
    result["bevoegd_gezag_nummer"] = vestigingen[0]["bevoegd_gezag_nummer"] if vestigingen else None

    if len(personeel_rows) == 0:
        result["status"] = "geen_personeelsregel"
        result["opmerkingen"].append("Geen personeelsregel gevonden voor dit BRIN in de personeelsbron.")
        return result

    if len(personeel_rows) > 1:
        result["status"] = "handmatige_controle"
        actueel_jaar = max(year_cols.keys()) if year_cols else None
        result["actueel_jaar"] = actueel_jaar
        details = []
        for row in personeel_rows:
            ow = row.get("ONDERWIJSTYPE")
            waarde = row.get(year_cols[actueel_jaar]) if actueel_jaar else None
            bg = normalize_bevoegd_gezag(row.get("BEVOEGD GEZAG"))
            details.append(f"ONDERWIJSTYPE={ow!r} BEVOEGD GEZAG={bg!r} PERSONEN {actueel_jaar}={waarde!r}")
        result["opmerkingen"].append(
            f"Meer dan 1 personeelsregel ({len(personeel_rows)}) gevonden voor dit BRIN; "
            "niet automatisch opgeteld (nog niet methodologisch vastgesteld dat sommeren "
            "veilig is). Geen segment bepaald. Gevonden regels: " + " || ".join(details)
        )
        if personeel_rows:
            result["bevoegd_gezag_nummer"] = normalize_bevoegd_gezag(personeel_rows[0].get("BEVOEGD GEZAG")) or result["bevoegd_gezag_nummer"]
        return result

    # Exact 1 personeelsregel: normale verwerking.
    result["status"] = "ok"
    row = personeel_rows[0]
    bg = normalize_bevoegd_gezag(row.get("BEVOEGD GEZAG"))
    if bg:
        result["bevoegd_gezag_nummer"] = bg

    waarde_info = bepaal_actuele_waarde(row, year_cols)
    result["actueel_jaar"] = waarde_info["actueel_jaar"]
    result["actuele_bronwaarde"] = waarde_info["actuele_bronwaarde"]
    result["aantal_personen"] = waarde_info["aantal_personen"]
    result["waarde_status"] = waarde_info["waarde_status"]
    result["peildatum"] = waarde_info["peildatum"]
    result["laatste_numerieke_waarde"] = waarde_info["laatste_numerieke_waarde"]
    result["laatste_numerieke_jaar"] = waarde_info["laatste_numerieke_jaar"]
    result["laatste_numerieke_peildatum"] = waarde_info["laatste_numerieke_peildatum"]
    result["segment"] = waarde_info["segment"]

    if waarde_info["waarde_status"] == "numeriek":
        result["bronniveau_medewerkers"] = "B"
        result["actualiteit_medewerkers"] = "Expliciet gedateerd"
        result["meetniveau_medewerkers"] = "Organisatie"
        result["eenheid_medewerkers"] = "Werkzame personen"
    else:
        result["opmerkingen"].append(
            f"Actuele waarde (PERSONEN {waarde_info['actueel_jaar']}) is niet-numeriek "
            f"({waarde_info['oorspronkelijke_waarde']!r}) en is NIET naar 0 omgezet. "
            "Aantal medewerkers = Onbekend, Segment = onbekend."
        )
        if waarde_info["laatste_numerieke_jaar"] is not None:
            result["opmerkingen"].append(
                f"Historische context (bepaalt segment NIET): laatste numerieke waarde "
                f"{waarde_info['laatste_numerieke_waarde']} in {waarde_info['laatste_numerieke_jaar']}."
            )

    return result


# ---------------------------------------------------------------------------
# Terminaloutput
# ---------------------------------------------------------------------------

def print_resultaat(result):
    print("=" * 70)
    print(f"BRIN: {result['brin']}")
    print(f"Status: {result['status']}")
    print(f"Sector: {result['sector']}")
    print(f"Instellingsnaam (indicatie): {result['instellingsnaam_indicatie']}")
    print(f"Bevoegd gezag nummer: {result['bevoegd_gezag_nummer']}")
    print(f"Aantal vestigingen: {result['aantal_vestigingen']}")
    for v in result["vestigingen"]:
        print(
            f"  - {v['vestigingscode']}: {v['vestigingsnaam']} | "
            f"{v['adres']}, {v['postcode']} {v['plaats']} | "
            f"bevoegd gezag {v['bevoegd_gezag_nummer']}"
        )
    print(f"Personeelsbron: {result['personeelsbron']}")
    print(f"Personeelsbron type: {result['personeelsbron_type']}")
    print(f"Actueel jaar: {result['actueel_jaar']}")
    print(f"Actuele bronwaarde: {result['actuele_bronwaarde']!r}")
    print(f"Aantal personen: {result['aantal_personen']}")
    print(f"Waarde status: {result['waarde_status']}")
    print(f"Peildatum: {result['peildatum']}")
    print(f"Laatste numerieke waarde (context): {result['laatste_numerieke_waarde']}")
    print(f"Laatste numerieke jaar (context): {result['laatste_numerieke_jaar']}")
    print(f"Laatste numerieke peildatum (context): {result['laatste_numerieke_peildatum']}")
    print(f"Bronniveau medewerkers: {result['bronniveau_medewerkers']}")
    print(f"Actualiteit medewerkers: {result['actualiteit_medewerkers']}")
    print(f"Meetniveau medewerkers: {result['meetniveau_medewerkers']}")
    print(f"Eenheid medewerkers: {result['eenheid_medewerkers']}")
    print(f"Segment: {result['segment']}")
    if result["opmerkingen"]:
        print("Opmerkingen:")
        for o in result["opmerkingen"]:
            print(f"  - {o}")
    print("=" * 70)


# ---------------------------------------------------------------------------
# Zelftests (python scripts/duo_onderwijs_lookup.py --selftest)
# ---------------------------------------------------------------------------

def _run_selftest():
    fouten = []

    def check(label, cond):
        status = "OK" if cond else "FOUT"
        print(f"[{status}] {label}")
        if not cond:
            fouten.append(label)

    print("TEST 1: BRIN = 17HN")
    r1 = lookup_brin("17HN")
    check("sector == VO", r1["sector"] == "VO")
    check("bevoegd_gezag_nummer == '41211'", r1["bevoegd_gezag_nummer"] == "41211")
    check("aantal_vestigingen == 4", r1["aantal_vestigingen"] == 4)
    check("status == ok (exact 1 personeelsregel)", r1["status"] == "ok")
    check("actuele_bronwaarde == 440", r1["actuele_bronwaarde"] == 440)
    check("aantal_personen == 440", r1["aantal_personen"] == 440)
    check("segment == 50+", r1["segment"] == "50+")
    check("peildatum == 01-10-2025", r1["peildatum"] == "01-10-2025")
    print()

    print("TEST 2: BRIN = 20GB")
    r2 = lookup_brin("20GB")
    check("sector == PO/speciaal onderwijs", r2["sector"] == "PO/speciaal onderwijs")
    check("bevoegd_gezag_nummer == '48348'", r2["bevoegd_gezag_nummer"] == "48348")
    check(
        "vestigingscode 20GB00 aanwezig",
        any(v["vestigingscode"] == "20GB00" for v in r2["vestigingen"]),
    )
    check("status == ok (exact 1 personeelsregel)", r2["status"] == "ok")
    check("actuele_bronwaarde == 18 (numeriek)", try_numeric(r2["actuele_bronwaarde"]) == 18)
    check("aantal_personen == 18", r2["aantal_personen"] == 18)
    check("segment == nog te bepalen", r2["segment"] == "nog te bepalen")
    check("peildatum == 01-10-2025", r2["peildatum"] == "01-10-2025")
    check("personeelsbron_type == afgeleide_compacte_DUO_dataset", r2["personeelsbron_type"] == "afgeleide_compacte_DUO_dataset")
    print()

    print("TEST 3: input '17hn' (kleine letters) == resultaat '17HN'")
    r3 = lookup_brin("17hn")
    check("resultaat identiek aan TEST 1", r3 == r1)
    print()

    print("TEST 4: input '17HN00' (vestigingscode) moet geweigerd worden")
    try:
        lookup_brin("17HN00")
        check("ValueError opgeworpen", False)
    except ValueError as e:
        msg = str(e)
        check("foutmelding noemt 'vestigingscode'", "vestigingscode" in msg.lower())
        check("foutmelding noemt '4 tekens'", "4 tekens" in msg)
    print()

    print("TEST 5: niet-bestaande code 'ZZZZ' mag niet crashen")
    try:
        r5 = lookup_brin("ZZZZ")
        check("geen crash", True)
        check("status == niet_gevonden", r5["status"] == "niet_gevonden")
    except Exception as e:
        check(f"geen crash (kreeg {type(e).__name__}: {e})", False)
    print()

    print("=" * 70)
    if fouten:
        print(f"ZELFTEST GEFAALD: {len(fouten)} check(s) niet geslaagd.")
        for f in fouten:
            print(f"  - {f}")
        return False
    print("ZELFTEST GESLAAGD: alle checks OK.")
    return True


def main():
    if len(sys.argv) != 2:
        print("Gebruik: python scripts/duo_onderwijs_lookup.py <BRIN>")
        print("     of: python scripts/duo_onderwijs_lookup.py --selftest")
        sys.exit(2)

    arg = sys.argv[1]
    if arg == "--selftest":
        ok = _run_selftest()
        sys.exit(0 if ok else 1)

    try:
        result = lookup_brin(arg)
    except ValueError as e:
        print(f"FOUT: {e}")
        sys.exit(1)

    print_resultaat(result)


if __name__ == "__main__":
    main()
